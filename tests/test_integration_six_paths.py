import hashlib
import shutil
from pathlib import Path

import pytest
from docx import Document

from local_doc_converter.converter import DocumentConverter


pytestmark = pytest.mark.skipif(shutil.which("pandoc") is None, reason="本机未安装 Pandoc")


def test_all_six_conversion_paths(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    fake_home = tmp_path / "home"
    fake_home.mkdir()
    monkeypatch.setenv("HOME", str(fake_home))
    inputs = fake_home / "inputs"
    inputs.mkdir()
    output = fake_home / "output"

    txt = inputs / "sample.txt"
    txt.write_text("第一章 示例\n\n正文内容。\n\n- 项目一\n- 项目二", encoding="utf-8")
    markdown = inputs / "sample.md"
    markdown.write_text("# 示例\n\n正文。\n\n| 名称 | 数量 |\n|---|---|\n| 苹果 | 2 |", encoding="utf-8")
    docx = inputs / "sample.docx"
    document = Document()
    document.add_heading("示例", level=1)
    document.add_paragraph("正文内容。")
    document.add_paragraph("列表项目", style="List Bullet")
    document.save(docx)

    converter = DocumentConverter()
    cases = [
        (txt, "markdown"), (txt, "docx"),
        (markdown, "txt"), (markdown, "docx"),
        (docx, "markdown"), (docx, "txt"),
    ]
    original_hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in (txt, markdown, docx)}
    results = [converter.convert(source, target, output) for source, target in cases]

    assert all(result.report.success for result in results), [result.report.error for result in results]
    assert all(result.output_path and result.output_path.stat().st_size > 0 for result in results)
    assert original_hashes == {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in original_hashes}


def test_markdown_soft_breaks_are_kept_in_txt(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    fake_home = tmp_path / "home"
    fake_home.mkdir()
    monkeypatch.setenv("HOME", str(fake_home))
    markdown = fake_home / "换行.md"
    markdown.write_text("# 标题\n\n第一行\n第二行\n\n下一段\n", encoding="utf-8")

    result = DocumentConverter().convert(markdown, "txt", fake_home / "output")

    assert result.report.success, result.report.error
    assert result.output_path is not None
    assert "第一行\n第二行\n\n下一段" in result.output_path.read_text(encoding="utf-8")
